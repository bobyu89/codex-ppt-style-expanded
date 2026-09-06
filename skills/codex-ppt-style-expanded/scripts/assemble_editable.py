"""Assemble native editable text and shapes over separately generated backgrounds."""
import argparse
import json
from pathlib import Path
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.parts.image import Image

def color(value):
    return RGBColor.from_string(value.lstrip('#'))

def box_coords(box, width, height):
    x, y, w, h = box
    if not all(isinstance(v, (int, float)) for v in box):
        raise ValueError('box coordinates must be numbers')
    if x < 0 or y < 0 or w <= 0 or h <= 0 or x+w > 1.000001 or y+h > 1.000001:
        raise ValueError('box is outside slide: ' + str(box))
    return Inches(x*width), Inches(y*height), Inches(w*width), Inches(h*height)

def add_picture(slide, path, box, width, height, background=False):
    if not path.is_file():
        raise FileNotFoundError(path)
    pxw, pxh = Image.from_file(str(path)).size
    x, y, w, h = box_coords(box, width, height)
    if background and abs((pxw/pxh)/(width/height)-1) > 0.02:
        raise ValueError('background aspect ratio differs from slide: ' + str(path))
    scale = min(w/pxw, h/pxh)
    pw, ph = int(pxw*scale), int(pxh*scale)
    slide.shapes.add_picture(str(path), int(x+(w-pw)/2), int(y+(h-ph)/2), pw, ph)

def assemble(spec_path, output):
    data = json.loads(spec_path.read_text(encoding='utf-8'))
    width, height = data.get('width', 13.333333), data.get('height', 7.5)
    if width <= 0 or height <= 0 or not data.get('slides'):
        raise ValueError('positive dimensions and at least one slide required')
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(width), Inches(height)
    for page in data['slides']:
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        fill = slide.background.fill
        fill.solid()
        fill.fore_color.rgb = color(page.get('background_color', 'FFFFFF'))
        if page.get('background'):
            add_picture(slide, spec_path.parent/page['background'], [0,0,1,1], width, height, True)
        for item in page.get('elements', []):
            coords = box_coords(item['box'], width, height)
            if item['type'] == 'image':
                add_picture(slide, spec_path.parent/item['path'], item['box'], width, height)
            elif item['type'] == 'rect':
                shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, *coords)
                shape.fill.solid()
                shape.fill.fore_color.rgb = color(item.get('fill', 'FFFFFF'))
                if item.get('line'):
                    shape.line.color.rgb = color(item['line'])
                else:
                    shape.line.fill.background()
            elif item['type'] == 'text':
                tf = slide.shapes.add_textbox(*coords).text_frame
                tf.word_wrap = True
                tf.margin_left = tf.margin_right = 0
                tf.margin_top = tf.margin_bottom = 0
                for index, line in enumerate(item['text'].split('\n')):
                    paragraph = tf.paragraphs[0] if index == 0 else tf.add_paragraph()
                    paragraph.alignment = {'left': PP_ALIGN.LEFT, 'center': PP_ALIGN.CENTER,
                                           'right': PP_ALIGN.RIGHT}[item.get('align', 'left')]
                    run = paragraph.add_run()
                    run.text = line
                    run.font.name = item.get('font', 'Microsoft JhengHei')
                    run.font.size = Pt(item.get('size', 24))
                    run.font.bold = item.get('bold', False)
                    run.font.color.rgb = color(item.get('color', '182B35'))
            else:
                raise ValueError('unsupported element type: ' + item['type'])
        slide.notes_slide.notes_text_frame.text = page.get('notes', '')
    output.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(output))

if __name__ == '__main__':
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('spec', type=Path)
    parser.add_argument('output', type=Path)
    args = parser.parse_args()
    assemble(args.spec.resolve(), args.output.resolve())
